"""M4: the full demo loop — corpus → model → opportunities → executive
artifacts, with citation coverage, numeric integrity, and immutability."""

import io
import re

import docx as docx_lib
from evals import corpus
from openpyxl import load_workbook
from pptx import Presentation

from tests.helpers import create_project, ingest_document

M4_CORPUS = [
    "claims-intake-sop.md",
    "adjuster-assignment-procedure.docx",
    "ops-weekly-2026-05-12.md",
]


def _corpus_doc(name: str) -> corpus.CorpusDoc:
    return next(d for d in corpus.build_corpus() if d.filename == name)


def _prepared_project(client, as_user):
    org_id, project_id = create_project(client, as_user, "M4")
    for name in M4_CORPUS:
        ingest_document(client, org_id, project_id, _corpus_doc(name))
    for path in ("knowledge/extract", "opportunities/generate"):
        run = client.post(f"/v1/orgs/{org_id}/projects/{project_id}/{path}")
        status = client.get(
            f"/v1/orgs/{org_id}/pipeline-runs/{run.json()['pipeline_run_id']}"
        ).json()
        assert status["status"] == "succeeded", status
    client.post(f"/v1/orgs/{org_id}/projects/{project_id}/roadmap/generate")
    return org_id, project_id


def _generate(client, org_id, project_id, deliverable_type):
    res = client.post(
        f"/v1/orgs/{org_id}/projects/{project_id}/deliverables",
        json={"type": deliverable_type},
    )
    assert res.status_code == 202, res.text
    body = res.json()
    run = client.get(
        f"/v1/orgs/{org_id}/pipeline-runs/{body['pipeline_run_id']}"
    ).json()
    assert run["status"] == "succeeded", run
    versions = client.get(
        f"/v1/orgs/{org_id}/deliverables/{body['deliverable']['id']}/versions"
    ).json()
    return body["deliverable"], versions[0]


def _download(client, org_id, version_id, fmt) -> bytes:
    res = client.get(
        f"/v1/orgs/{org_id}/deliverable-versions/{version_id}/download",
        params={"format": fmt},
    )
    assert res.status_code == 200, res.text
    return res.content


def test_all_four_deliverable_types_generate_with_citations(client, as_user):
    org_id, project_id = _prepared_project(client, as_user)
    for deliverable_type in (
        "executive_summary",
        "current_state_assessment",
        "opportunity_register",
        "roadmap_deck",
    ):
        _, version = _generate(client, org_id, project_id, deliverable_type)
        assert version["version"] == 1
        assert version["formats"], deliverable_type
        if deliverable_type != "roadmap_deck":
            assert version["citation_count"] > 0, (
                f"{deliverable_type} must carry a citation appendix"
            )


def test_pptx_deck_structure_and_speaker_note_citations(client, as_user):
    org_id, project_id = _prepared_project(client, as_user)
    _, version = _generate(client, org_id, project_id, "executive_summary")
    data = _download(client, org_id, version["id"], "pptx")

    deck = Presentation(io.BytesIO(data))
    all_text = "\n".join(
        shape.text_frame.text
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert len(deck.slides) >= 5  # title, overview, findings, opps, roadmap, sources
    assert "Executive Summary" in all_text
    assert "Sources" in all_text
    notes = "\n".join(
        slide.notes_slide.notes_text_frame.text
        for slide in deck.slides
        if slide.has_notes_slide
    )
    assert "Sources:" in notes, "speaker notes must carry the slide's citations"


def test_docx_numeric_integrity_against_roi_sheets(client, as_user):
    """Every dollar figure in the export must exist in an ROI sheet or be a
    sum of them — no invented numbers (docs/01 §2.5)."""
    org_id, project_id = _prepared_project(client, as_user)
    opportunities = client.get(
        f"/v1/orgs/{org_id}/projects/{project_id}/opportunities"
    ).json()
    allowed = set()
    for opportunity in opportunities:
        computed = opportunity["roi"]["computed"]
        for key in ("gross_annual_savings", "net_annual_savings", "three_year_net"):
            if computed.get(key) is not None:
                allowed.add(round(computed[key]))
    allowed.add(
        round(sum(o["roi"]["computed"]["net_annual_savings"] for o in opportunities))
    )

    _, version = _generate(client, org_id, project_id, "executive_summary")
    data = _download(client, org_id, version["id"], "docx")
    document = docx_lib.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs) + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    for match in re.finditer(r"\$([\d,]+)\b", text):
        value = int(match.group(1).replace(",", ""))
        assert value in allowed, (
            f"figure ${value:,} in export does not trace to any ROI sheet"
        )


def test_xlsx_register_round_trip(client, as_user):
    org_id, project_id = _prepared_project(client, as_user)
    _, version = _generate(client, org_id, project_id, "opportunity_register")
    data = _download(client, org_id, version["id"], "xlsx")
    workbook = load_workbook(io.BytesIO(data))
    assert {"Register", "ROI assumptions", "Sources"} <= set(workbook.sheetnames)
    register = workbook["Register"]
    assert register.max_row >= 2, "register must contain opportunities"
    header = [c.value for c in register[1]]
    assert "Net annual savings" in header
    sources = workbook["Sources"]
    assert sources.max_row >= 2, "sources sheet must list citations"


def test_versions_are_immutable_and_increment(client, as_user):
    org_id, project_id = _prepared_project(client, as_user)
    deliverable, v1 = _generate(client, org_id, project_id, "executive_summary")
    v1_bytes = _download(client, org_id, v1["id"], "md")

    # Change the model (confirm an opportunity), regenerate → version 2.
    opportunities = client.get(
        f"/v1/orgs/{org_id}/projects/{project_id}/opportunities"
    ).json()
    client.post(
        f"/v1/orgs/{org_id}/opportunities/{opportunities[0]['id']}/review",
        json={"action": "confirm"},
    )
    _, v2 = _generate(client, org_id, project_id, "executive_summary")
    assert v2["version"] == 2

    versions = client.get(
        f"/v1/orgs/{org_id}/deliverables/{deliverable['id']}/versions"
    ).json()
    assert [v["version"] for v in versions] == [2, 1]

    # v1's artifact is byte-identical after v2 exists: immutability.
    assert _download(client, org_id, v1["id"], "md") == v1_bytes
    # And v2 differs (review coverage moved).
    assert _download(client, org_id, v2["id"], "md") != v1_bytes


def test_unavailable_format_is_a_clean_error(client, as_user):
    org_id, project_id = _prepared_project(client, as_user)
    _, version = _generate(client, org_id, project_id, "current_state_assessment")
    res = client.get(
        f"/v1/orgs/{org_id}/deliverable-versions/{version['id']}/download",
        params={"format": "pptx"},
    )
    assert res.status_code == 400


def test_spec_endpoint_hides_internal_keys(client, as_user):
    org_id, project_id = _prepared_project(client, as_user)
    _, version = _generate(client, org_id, project_id, "executive_summary")
    spec = client.get(
        f"/v1/orgs/{org_id}/deliverable-versions/{version['id']}/spec"
    ).json()["spec"]
    assert "_run_id" not in spec
    assert spec["sections"]
    stats = spec["sections"][0]["stats"]
    assert any(s["label"] == "Model analyst-reviewed" for s in stats)
