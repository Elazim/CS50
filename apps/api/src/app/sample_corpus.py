"""Synthetic insurance-claims corpus v1 (docs/03 §7, docs/06 M1).

A fictional mid-size P&C carrier ("Northbridge Mutual") with a claims
operation full of deliberate, realistic problems: manual rekeying between
systems, email-driven handoffs, unclear escalation, aging spreadsheets.
Used by demos, integration tests, and the eval suite — no environment ever
needs real customer data (docs/05 §1).

Binary formats (DOCX/XLSX/PPTX/PDF) are generated programmatically so the
repo stays text-only and the format writers are exercised too.
"""

import io
from dataclasses import dataclass

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@dataclass
class CorpusDoc:
    filename: str
    mime: str
    data: bytes


INTAKE_SOP = """\
# Standard Operating Procedure: First Notice of Loss (FNOL) Intake

## 1. Purpose
This standard operating procedure describes how the Claims Intake team receives,
registers, and routes new loss notifications at Northbridge Mutual.

## 2. Receiving a notification
New claims arrive through three channels: the customer portal, the call center,
and broker email. Portal claims land in ClaimCore automatically. Call center
agents capture details in CallTrak and then re-enter the same information into
ClaimCore manually. Broker emails arrive in the shared claims-intake mailbox and
are printed and keyed into ClaimCore by an intake coordinator.

## 3. Registration steps
- Verify the policy number in PolicyHub and confirm coverage is in force.
- Create the claim record in ClaimCore with loss date, loss type, and reporter details.
- Attach supporting documents received by email to the claim record.
- Set the initial severity using the triage guidelines.
- Route the claim to the assignment queue.

## 4. Known issues
Duplicate entry between CallTrak and ClaimCore is the team's biggest time sink and
a frequent source of transcription errors. Broker email intake regularly exceeds
the 24-hour registration target during storm season.
"""

ASSIGNMENT_SOP_SECTIONS = [
    ("Adjuster Assignment Procedure", 1, None),
    ("Assignment rules", 2, (
        "Registered claims wait in the assignment queue until a supervisor reviews them "
        "each morning. The supervisor assigns each claim to an adjuster based on loss "
        "type, severity, adjuster licensing state, and current workload. Workload is "
        "tracked in a shared spreadsheet that adjusters update themselves on Fridays."
    )),
    ("Severity routing", 2, "TABLE"),
    ("Escalation", 2, (
        "Claims with reserves above $250,000, suspected fraud indicators, or attorney "
        "involvement are escalated to the Complex Claims unit. Escalation happens by "
        "email to the unit manager; there is no queue or tracking of escalated claims."
    )),
    ("Reassignment", 2, (
        "When an adjuster is out for more than three days, their supervisor manually "
        "reviews open claims and reassigns urgent ones. Non-urgent claims wait for the "
        "adjuster to return."
    )),
]

ASSIGNMENT_TABLE = [
    ["Severity", "Loss types", "Assigned to", "Target"],
    ["1 - Minor", "Glass, towing, small property", "Express team", "Same day"],
    ["2 - Standard", "Auto collision, water damage", "Field adjusters", "1 business day"],
    ["3 - Major", "Fire, injury, total loss", "Senior adjusters", "4 business hours"],
    ["4 - Complex", "Litigation, fraud referral, catastrophe", "Complex Claims unit", "2 hours"],
]

TRIAGE_POLICY = """\
# Claims Triage and Severity Policy

## Policy statement
This policy defines the severity tiers used to triage incoming claims and the
handling guidelines for each tier. All registered claims must receive a severity
rating before assignment.

## Severity tiers
Severity 1 covers minor losses under $5,000 with no injuries, such as glass and
towing. Severity 2 covers standard losses between $5,000 and $50,000. Severity 3
covers major losses above $50,000 or any claim involving bodily injury. Severity 4
covers complex matters: suspected fraud, litigation, catastrophe events, and any
reserve above $250,000.

## Compliance requirement
State fair-claims-handling regulations require acknowledgement of every claim
within 24 hours of receipt and a coverage decision within 30 days. Severity
ratings must be reviewed by a supervisor within one business day.
"""

SUBROGATION = """\
# Subrogation Referral Process

## Identifying recovery potential
Adjusters flag claims with subrogation potential when a third party may be liable
for the loss. The flag is a free-text note in ClaimCore; there is no structured
field, so the recovery team runs a weekly keyword search over claim notes to find
candidates and misses an estimated one in five referrals.

## Referral steps
- Adjuster documents third-party liability evidence in the claim file.
- Recovery coordinator opens a subrogation file in a separate Access database.
- Demand letters are generated in Word from a template and tracked in a spreadsheet.
- Recoveries are posted back to ClaimCore manually at month end.
"""

MEETING_NOTES = """\
# Claims Operations Weekly — Meeting Notes, May 12 2026

Attendees: M. Okafor (Claims Director), J. Reyes (Intake Supervisor),
P. Lindqvist (IT Liaison), S. Boucher (Complex Claims Manager)

## Discussion
J. Reyes reported that intake coordinators spend roughly a third of their day
re-entering CallTrak notifications into ClaimCore, and error rates spike every
storm surge. P. Lindqvist confirmed the CallTrak vendor exposes an API that has
never been integrated. S. Boucher raised that escalations by email keep getting
lost; two complex claims last month sat unassigned for over a week.

## Decisions
The team agreed to scope an integration between CallTrak and ClaimCore this
quarter and to pilot a shared escalation queue for the Complex Claims unit.

## Action items
- P. Lindqvist: request CallTrak API documentation from the vendor by May 19.
- J. Reyes: measure intake rekeying time for one week using the time-tracking template.
- S. Boucher: draft requirements for the escalation queue pilot.
"""

GLOSSARY = """\
Claims Glossary — Northbridge Mutual

FNOL: First Notice of Loss, the initial report of a claim.
ClaimCore: the core claims administration system of record.
CallTrak: the call center telephony and note-taking application.
PolicyHub: the policy administration system used to verify coverage.
Reserve: the amount set aside to pay a claim's expected cost.
Subrogation: recovering claim costs from a liable third party.
Complex Claims unit: the team handling litigation, fraud and catastrophe claims.
"""

FRAUD_SECTIONS = [
    ("Fraud Referral Procedure", 1, None),
    ("Indicators", 2, (
        "Adjusters watch for common fraud indicators: losses reported shortly after "
        "policy inception, prior similar claims, inconsistent statements, pressure for "
        "quick settlement, and invoices from unknown vendors."
    )),
    ("Referral", 2, (
        "Suspected fraud is referred to the Special Investigations Unit by completing "
        "the SIU referral form, a fillable PDF emailed to the SIU inbox. The adjuster "
        "pauses payment activity and documents the indicators in the claim notes. SIU "
        "aims to acknowledge referrals within two business days but does not report "
        "acknowledgement times today."
    )),
    ("Regulatory obligations", 2, (
        "Confirmed fraud must be reported to the state fraud bureau within 30 days. "
        "The compliance team files these reports manually from a quarterly extract."
    )),
]

COMPLAINTS_ROWS = [
    ["Date", "Policy #", "Claim #", "Contact email", "Contact phone", "Summary"],
    ["2026-04-02", "POL-4482911", "CLM-2201435", "rita.moreno@example.com",
     "555-201-8890", "No update on water damage claim for 12 days"],
    ["2026-04-09", "POL-9034522", "CLM-2201619", "d.chen@example.net",
     "555-318-4471", "Asked to resubmit invoices already sent by email"],
    ["2026-04-15", "POL-1120087", "CLM-2202044", "kwame.a@example.org",
     "555-742-0913", "Adjuster changed twice; had to re-explain the loss"],
    ["2026-04-23", "POL-5566190", "CLM-2202307", "l.bianchi@example.com",
     "555-609-2284", "Settlement letter referenced the wrong vehicle"],
]

VOLUME_ROWS = [
    ["Month", "Claims received", "Avg registration minutes", "Avg days to assignment",
     "Rekeyed from CallTrak"],
    ["2026-01", "1840", "22", "1.8", "1102"],
    ["2026-02", "1710", "21", "1.6", "1034"],
    ["2026-03", "2260", "26", "2.4", "1391"],
    ["2026-04", "1985", "24", "2.1", "1187"],
]

ROSTER_ROWS = [
    ["Team", "Role", "Headcount", "Location"],
    ["Claims Intake", "Intake Coordinator", "9", "Hartford"],
    ["Claims Intake", "Intake Supervisor", "1", "Hartford"],
    ["Field Adjusting", "Adjuster", "24", "Remote"],
    ["Field Adjusting", "Senior Adjuster", "8", "Remote"],
    ["Complex Claims", "Complex Adjuster", "6", "Hartford"],
    ["Recovery", "Subrogation Coordinator", "3", "Hartford"],
]

KICKOFF_SLIDES = [
    ("Claims Transformation Initiative", "Kickoff — June 2026",
     "Welcome. Goal today is alignment on scope and the first ninety days."),
    ("Why now", "Intake rekeying consumes 30% of coordinator time. "
     "Escalations by email get lost. Assignment waits for a daily manual review.",
     "These three pain points came directly from the ops weekly and the volume report."),
    ("First 90 days", "Integrate CallTrak with ClaimCore. Pilot the escalation queue. "
     "Automate assignment for severity 1 and 2 claims.",
     "Quick wins first; the assignment automation depends on the workload data being live."),
]

VENDOR_INVOICE_PDF = [
    ("Vendor Invoice Approval Procedure", 1),
    ("Scope", 2),
    ("This procedure covers approval of vendor invoices for claim-related services "
     "such as independent appraisals, towing, and restoration work.", 0),
    ("Approval thresholds", 2),
    ("Invoices under $2,500 are approved by the handling adjuster. Invoices between "
     "$2,500 and $25,000 require supervisor approval. Larger invoices require the "
     "claims director's signature, collected by routing a paper form.", 0),
    ("Payment", 2),
    ("Approved invoices are entered into the finance system by the accounts payable "
     "team from scanned copies, typically within five business days.", 0),
]


def _docx(sections: list[tuple], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    doc = Document()
    for title, level, body in sections:
        if level:
            doc.add_heading(title, level=level)
        if body == "TABLE" and table:
            t = doc.add_table(rows=0, cols=len(table[0]))
            for row in table:
                cells = t.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = value
        elif body:
            doc.add_paragraph(body)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx(sheet_name: str, rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx(slides: list[tuple[str, str, str]]) -> bytes:
    from pptx import Presentation

    deck = Presentation()
    layout = deck.slide_layouts[1]  # title + content
    for title, body, notes in slides:
        slide = deck.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def _pdf(blocks: list[tuple[str, int]]) -> bytes:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    y = 72.0
    for text, level in blocks:
        size = 18 if level == 1 else 14 if level == 2 else 10
        font = "hebo" if level else "helv"
        rect = fitz.Rect(72, y, page.rect.width - 72, page.rect.height - 72)
        used = page.insert_textbox(rect, text, fontsize=size, fontname=font)
        y += (rect.height - used) + (10 if level else 6)
        if y > page.rect.height - 144:
            page = doc.new_page()
            y = 72.0
    data = doc.tobytes()
    doc.close()
    return data


def build_corpus() -> list[CorpusDoc]:
    return [
        CorpusDoc("claims-intake-sop.md", "text/markdown", INTAKE_SOP.encode()),
        CorpusDoc(
            "adjuster-assignment-procedure.docx", DOCX_MIME,
            _docx(ASSIGNMENT_SOP_SECTIONS, ASSIGNMENT_TABLE),
        ),
        CorpusDoc("claims-triage-policy.md", "text/markdown", TRIAGE_POLICY.encode()),
        CorpusDoc("subrogation-process.md", "text/markdown", SUBROGATION.encode()),
        CorpusDoc("ops-weekly-2026-05-12.md", "text/markdown", MEETING_NOTES.encode()),
        CorpusDoc("claims-glossary.txt", "text/plain", GLOSSARY.encode()),
        CorpusDoc("fraud-referral-procedure.docx", DOCX_MIME, _docx(FRAUD_SECTIONS)),
        CorpusDoc("customer-complaints-log.xlsx", XLSX_MIME,
                  _xlsx("Complaints", COMPLAINTS_ROWS)),
        CorpusDoc("claims-volume-report.xlsx", XLSX_MIME, _xlsx("Volumes", VOLUME_ROWS)),
        CorpusDoc("org-roster.xlsx", XLSX_MIME, _xlsx("Roster", ROSTER_ROWS)),
        CorpusDoc("transformation-kickoff.pptx", PPTX_MIME, _pptx(KICKOFF_SLIDES)),
        CorpusDoc("vendor-invoice-approval.pdf", "application/pdf",
                  _pdf(VENDOR_INVOICE_PDF)),
    ]
