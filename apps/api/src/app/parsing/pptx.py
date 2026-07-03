"""PPTX via python-pptx: slide titles become headings, body frames become
paragraphs, speaker notes are kept (they often carry the real process
knowledge in operations decks)."""

import io

from pptx import Presentation

from app.modules.documents.elements import ElementKind
from app.parsing.base import ParsedElement, UnsupportedDocumentError


def parse_pptx(data: bytes) -> list[ParsedElement]:
    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise UnsupportedDocumentError(f"Cannot open PPTX: {exc}") from exc

    elements: list[ParsedElement] = []
    for slide_no, slide in enumerate(deck.slides, start=1):
        title = ""
        if slide.shapes.title and slide.shapes.title.text.strip():
            title = " ".join(slide.shapes.title.text.split())
        section = [f"Slide {slide_no}" + (f": {title}" if title else "")]

        if title:
            elements.append(
                ParsedElement(
                    kind=ElementKind.heading,
                    text=title,
                    page=slide_no,
                    section_path=section,
                    meta={"level": 1},
                )
            )
        for shape in slide.shapes:
            if shape == slide.shapes.title or not shape.has_text_frame:
                continue
            text = " ".join(shape.text_frame.text.split())
            if text:
                elements.append(
                    ParsedElement(
                        kind=ElementKind.paragraph,
                        text=text,
                        page=slide_no,
                        section_path=section,
                    )
                )
        if slide.has_notes_slide:
            notes = " ".join(slide.notes_slide.notes_text_frame.text.split())
            if notes:
                elements.append(
                    ParsedElement(
                        kind=ElementKind.note,
                        text=notes,
                        page=slide_no,
                        section_path=section,
                        meta={"speaker_notes": True},
                    )
                )
    return elements
