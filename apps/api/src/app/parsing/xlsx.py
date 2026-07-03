"""XLSX via openpyxl: each sheet becomes a table element (rows capped);
tables stay structured — flattening spreadsheets to prose destroys them."""

import io

from openpyxl import load_workbook

from app.modules.documents.elements import ElementKind
from app.parsing.base import ParsedElement, UnsupportedDocumentError

MAX_ROWS_PER_SHEET = 500


def parse_xlsx(data: bytes) -> list[ParsedElement]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise UnsupportedDocumentError(f"Cannot open XLSX: {exc}") from exc

    elements: list[ParsedElement] = []
    for sheet in workbook.worksheets:
        rows: list[list[str]] = []
        truncated = False
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= MAX_ROWS_PER_SHEET:
                truncated = True
                break
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
        elements.append(
            ParsedElement(
                kind=ElementKind.table,
                text="\n".join(" | ".join(row) for row in rows),
                section_path=[sheet.title],
                table_json={"rows": rows, "truncated": truncated},
                meta={"sheet": sheet.title},
            )
        )
    return elements
