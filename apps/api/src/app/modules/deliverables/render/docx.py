"""DOCX renderer via python-docx."""

import io

import docx
from docx.shared import Pt, RGBColor

ACCENT = RGBColor(0x5A, 0x56, 0xC9)
MUTED = RGBColor(0x5C, 0x65, 0x70)


def _marks(item: dict) -> str:
    marks = "".join(f"[{n}]" for n in item.get("citations", []))
    if item.get("unreviewed"):
        marks += " *"
    return f" {marks}" if marks else ""


def render_docx(spec: dict) -> bytes:
    document = docx.Document()

    title = document.add_heading(spec["title"], level=0)
    title.runs[0].font.color.rgb = ACCENT
    meta = document.add_paragraph(
        f"Generated {spec['generated_at'][:10]} · model {spec['review_coverage']} "
        "analyst-reviewed · figures derive from the ROI assumption sheets"
    )
    meta.runs[0].font.size = Pt(9)
    meta.runs[0].font.color.rgb = MUTED

    for section in spec["sections"]:
        document.add_heading(section["heading"], level=1)
        kind = section["kind"]
        if kind == "overview":
            document.add_paragraph(section["body"])
            table = document.add_table(rows=0, cols=2)
            table.style = "Light Grid Accent 1"
            for stat in section["stats"]:
                cells = table.add_row().cells
                cells[0].text = stat["label"]
                cells[1].text = stat["value"]
        elif kind == "findings":
            for item in section["items"]:
                text = item["text"] + _marks(item)
                if item.get("taxonomy"):
                    text += f"  ({item['taxonomy']})"
                document.add_paragraph(text, style="List Bullet")
        elif kind == "opportunities":
            table = document.add_table(rows=1, cols=7)
            table.style = "Light Grid Accent 1"
            for i, header in enumerate(
                ["Opportunity", "Type", "Impact", "Complexity", "Risk",
                 "Net annual", "Payback"]
            ):
                table.rows[0].cells[i].text = header
            for row in section["rows"]:
                cells = table.add_row().cells
                cells[0].text = row["title"] + _marks(row)
                cells[1].text = row["taxonomy"]
                cells[2].text = str(row["impact"])
                cells[3].text = str(row["complexity"])
                cells[4].text = str(row["risk"])
                cells[5].text = (
                    f"${row['net_annual_savings']:,.0f}"
                    if row.get("net_annual_savings") is not None else "—"
                )
                cells[6].text = (
                    f"{row['payback_months']} mo" if row.get("payback_months") else "—"
                )
        elif kind == "roadmap":
            for horizon, label in (
                ("30", "First 30 days"), ("90", "90 days"),
                ("180", "6 months"), ("365", "12 months"),
            ):
                document.add_heading(label, level=2)
                items = section["horizons"].get(horizon, [])
                for item in items:
                    document.add_paragraph(
                        item["title"] + (" (quick win)" if item.get("quick_win") else ""),
                        style="List Bullet",
                    )
                if not items:
                    document.add_paragraph("—")
        elif kind == "process":
            table = document.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            for i, header in enumerate(["#", "Step", "Type", "Performer"]):
                table.rows[0].cells[i].text = header
            for step in section["steps"]:
                cells = table.add_row().cells
                cells[0].text = str(step["order"])
                cells[1].text = step["name"] + _marks(step)
                cells[2].text = step["step_type"]
                cells[3].text = step["performer"]
            if section["pain_points"]:
                document.add_heading("Pain points", level=2)
                for pain in section["pain_points"]:
                    text = pain["text"] + _marks(pain)
                    if pain.get("taxonomy"):
                        text += f"  ({pain['taxonomy']})"
                    document.add_paragraph(text, style="List Bullet")
        elif kind == "table":
            table = document.add_table(rows=1, cols=len(section["columns"]))
            table.style = "Light Grid Accent 1"
            for i, header in enumerate(section["columns"]):
                table.rows[0].cells[i].text = header
            for row in section["rows"]:
                cells = table.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = str(value)

    if spec["citations"]:
        document.add_heading("Sources", level=1)
        for citation in spec["citations"]:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(f"[{citation['n']}] {citation['filename']}")
            run.bold = True
            if citation.get("quote"):
                quote_run = paragraph.add_run(f' — "{citation["quote"]}"')
                quote_run.font.color.rgb = MUTED
                quote_run.font.size = Pt(9)
    footnote = document.add_paragraph(
        "* item generated by AI and not yet analyst-reviewed"
    )
    footnote.runs[0].font.size = Pt(8)
    footnote.runs[0].font.color.rgb = MUTED

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
