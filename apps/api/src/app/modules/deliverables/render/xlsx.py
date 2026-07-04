"""XLSX renderer via openpyxl (opportunity register: the working backlog)."""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="5A56C9")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _sheet(workbook, title, headers):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    return sheet


def render_xlsx(spec: dict) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    for section in spec["sections"]:
        if section["kind"] == "opportunities":
            sheet = _sheet(
                workbook, "Register",
                ["Opportunity", "Type", "Impact", "Complexity", "Risk",
                 "Net annual savings", "Payback (months)", "3-year net",
                 "Reviewed", "Sources", "Rationale"],
            )
            for row in section["rows"]:
                sheet.append([
                    row["title"],
                    row["taxonomy"],
                    row["impact"],
                    row["complexity"],
                    row["risk"],
                    row.get("net_annual_savings"),
                    row.get("payback_months"),
                    row.get("three_year_net"),
                    "no" if row.get("unreviewed") else "yes",
                    ", ".join(str(n) for n in row.get("citations", [])),
                    row.get("rationale", ""),
                ])
            sheet.column_dimensions["A"].width = 60
            sheet.column_dimensions["K"].width = 80
        elif section["kind"] == "table":
            sheet = _sheet(workbook, section["heading"][:28], section["columns"])
            for row in section["rows"]:
                sheet.append(list(row))
            sheet.column_dimensions["A"].width = 50

    if spec["citations"]:
        sheet = _sheet(workbook, "Sources", ["#", "Document", "Quote"])
        for citation in spec["citations"]:
            sheet.append([citation["n"], citation["filename"], citation.get("quote")])
        sheet.column_dimensions["B"].width = 40
        sheet.column_dimensions["C"].width = 100

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
